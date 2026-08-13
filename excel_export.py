from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from models import LawsuitCase
from constants import status_to_name

FIELD_MAP = {
    "id": "ID",
    "case_no": "案件编号",
    "court_case_no": "案号",
    "document_type": "文书类型",
    "plaintiff": "原告（申请人）",
    "defendant": "被告（被申请人）",
    "case_type_name": "业务类型",
    "standard_cause_name": "标准案由",
    "court_name": "受理法院/仲裁委",
    "current_status": "案件状态",
    "assigned_contact": "对接人",
    "case_description": "案由描述",
    "claim_amount": "涉案金额",
    "deadline_date": "期限日期",
    "deadline_type": "期限类型",
    "remark": "备注",
    "created_by": "创建人",
    "created_at": "创建时间",
}


def export_to_excel(cases: List[LawsuitCase], filepath: str = "cases_export.xlsx") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "案件数据"

    keys = list(FIELD_MAP.keys())
    headers = list(FIELD_MAP.values())

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    even_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

    for row_idx, case in enumerate(cases, 2):
        for col_idx, key in enumerate(keys, 1):
            value = getattr(case, key, None)
            if key == "current_status" and value is not None:
                value = status_to_name(value)
            if key == "claim_amount" and value is not None:
                value = float(value)
            if key in ("created_at", "updated_at") and value is not None:
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_idx % 2 == 0:
                cell.fill = even_fill

    for col_idx, header in enumerate(headers, 1):
        max_length = len(str(header))
        for row_idx in range(2, len(cases) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value is not None:
                max_length = max(max_length, len(str(cell_value)))
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    ws.freeze_panes = "A2"
    wb.save(filepath)
    return filepath
